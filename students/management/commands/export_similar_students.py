"""
تصدير ملف Excel شامل للكشف عن الأسماء المتطابقة والمتشابهة وحساباتها المالية (للمراجعة اليدوية فقط، لا يعدّل شيئاً).

الأوراق في الملف:
  ملخص               : أرقام عامة وشرح الألوان
  متطابقة            : طلاب بنفس الاسم (بعد توحيد الهمزات/التاء المربوطة/الياء/المسافات)
  متشابهة            : أسماء متقاربة (أخطاء إملائية) أو اسم ناقص مقابل اسم أطول (ينقصه جزء من النسب)
  رقم القيد المكرر    : طلاب يتشاركون نفس رقم القيد
  المعاملات المالية   : كل إيصالات (مدفوعات) الطلاب الواردين في الأوراق السابقة
  الأقساط            : كل أقساط الطلاب الواردين في الأوراق السابقة

الاستخدام (من مجلد المشروع):
  python manage.py export_similar_students                    # كل أنواع الدراسة، السنة النشطة
  python manage.py export_similar_students --workers-only     # العمال فقط
  python manage.py export_similar_students --year-id 3        # سنة محددة
  python manage.py export_similar_students --threshold 0.85   # حساسية التشابه (أقل = نتائج أكثر)
  python manage.py export_similar_students --out /home/sameh123/x.xlsx

يحتاج مكتبة openpyxl:  pip install openpyxl
"""
import difflib
import os
import re
import time
from collections import defaultdict
from decimal import Decimal

from django.core.management.base import BaseCommand, CommandError

# ======================================================================
# دوال نقية (لا تعتمد على قاعدة البيانات) لتسهيل الاختبار
# ======================================================================
_TASH = re.compile('[\u064B-\u065F\u0670\u0640]')
_MAP = str.maketrans({'أ': 'ا', 'إ': 'ا', 'آ': 'ا', 'ة': 'ه', 'ى': 'ي', 'ؤ': 'و', 'ئ': 'ي'})


def norm(text):
    t = _TASH.sub('', text or '').translate(_MAP)
    return re.sub(r'\s+', ' ', t).strip()


def full_name(s):
    return f"{s.first_name or ''} {s.last_name or ''}".strip()


def compact(text):
    return norm(text).replace(' ', '')


def exact_groups(students):
    g = defaultdict(list)
    for s in students:
        k = compact(full_name(s))
        if k:
            g[k].append(s)
    out = []
    for members in g.values():
        if len(members) > 1:
            same_raw = len({re.sub(r'\s+', ' ', full_name(m)) for m in members}) == 1
            reason = 'اسم مطابق تماماً' if same_raw else 'مطابق بعد توحيد الحروف/المسافات'
            out.append((reason, sorted(members, key=lambda m: m.pk)))
    return sorted(out, key=lambda x: norm(full_name(x[1][0])))


def similar_pairs(students, threshold=0.88):
    """أزواج الأسماء المتقاربة: (سبب، طلاب الاسم الأول + طلاب الاسم الثاني)."""
    by_key = defaultdict(list)
    for s in students:
        toks = norm(full_name(s)).split()
        if toks:
            by_key[''.join(toks)].append(s)
    toks_of = {k: norm(full_name(v[0])).split() for k, v in by_key.items()}
    keys = sorted(by_key, key=len)
    found = {}

    # 1) اسم أقصر هو بداية الاسم الأطول (ينقص أحدهما جزء من النسب)
    buckets = defaultdict(list)
    for k in keys:
        if len(toks_of[k]) >= 3:
            buckets[tuple(toks_of[k][:3])].append(k)
    for lst in buckets.values():
        for a in lst:
            for b in lst:
                ta, tb = toks_of[a], toks_of[b]
                if len(ta) < len(tb) <= len(ta) + 3 and tb[:len(ta)] == ta:
                    found[(a, b)] = ('أحد الاسمين جزء من الآخر (ينقصه جزء من النسب)', 1.0)

    # 2) تقارب حروف (أخطاء إملائية)
    sm = difflib.SequenceMatcher(autojunk=False)
    for i, a in enumerate(keys):
        sm.set_seq2(a)
        for b in keys[i + 1:]:
            if len(b) - len(a) > 3:
                break
            sm.set_seq1(b)
            if sm.real_quick_ratio() < threshold or sm.quick_ratio() < threshold:
                continue
            r = sm.ratio()
            if r >= threshold and (a, b) not in found and (b, a) not in found:
                found[(a, b)] = (f'تقارب حروف {r:.0%}', r)

    out = []
    for (a, b), (reason, score) in found.items():
        out.append((reason, sorted(by_key[a] + by_key[b], key=lambda m: (norm(full_name(m)), m.pk))))
    return sorted(out, key=lambda x: norm(full_name(x[1][0])))


def registration_groups(students):
    g = defaultdict(list)
    for s in students:
        k = (s.registration_number or '').strip()
        if k:
            g[k].append(s)
    return [(f'نفس رقم القيد: {k}', sorted(v, key=lambda m: m.pk)) for k, v in sorted(g.items()) if len(v) > 1]


def clean_nid(n):
    n = (str(n) if n is not None else '').strip()
    return n[:-2] if n.endswith('.0') else n


def valid_nid(n):
    """رقم قومي مصري صالح شكلاً: 14 رقماً، يبدأ بـ 2 أو 3، وشهر/يوم ميلاد منطقيان."""
    n = clean_nid(n)
    return bool(re.fullmatch(r'[23]\d{13}', n) and 1 <= int(n[3:5]) <= 12 and 1 <= int(n[5:7]) <= 31)


def _pair_signal(a, b):
    na, nb = clean_nid(a.national_id), clean_nid(b.national_id)
    va, vb = valid_nid(na), valid_nid(nb)
    ma, mb = compact(a.mother_name), compact(b.mother_name)
    pa, pb = (a.phone or '').strip(), (b.phone or '').strip()
    same_mother = bool(ma and mb and ma == mb)
    diff_mother = bool(ma and mb and ma != mb)
    same_phone = bool(pa and pb and pa == pb)
    diff_dob = bool(a.date_of_birth and b.date_of_birth and a.date_of_birth != b.date_of_birth)
    if va and vb and na == nb:
        return 'SAME'
    if va != vb:                       # أحدهما برقم قومي غير صالح (وهمي) والآخر صالح = سجل مؤقت
        return 'TEMP'
    if va and vb:                      # كلاهما صالح ومختلف
        if same_mother or same_phone:
            return 'SIB'               # إخوة/توائم: نفس الأم والقومي مختلف
        if diff_mother or diff_dob:
            return 'DIFF'
        return 'REVIEW'
    if same_mother or same_phone:      # كلاهما غير صالح لكن الأم/الهاتف متطابق
        return 'LIKELY'
    return 'REVIEW'


def group_verdict(members):
    """تقييم آلي مبدئي للمجموعة (للاسترشاد فقط، والقرار النهائي لك)."""
    sig = {_pair_signal(members[i], members[j]) for i in range(len(members)) for j in range(i + 1, len(members))}
    if 'SAME' in sig:
        return 'SAME', 'نفس الشخص على الأرجح (رقم قومي متطابق)'
    if 'TEMP' in sig:
        return 'TEMP', 'سجل مؤقت (قومي غير صالح) مقابل سجل حقيقي: على الأرجح مكرر'
    if 'LIKELY' in sig:
        return 'LIKELY', 'يُرجَّح نفس الشخص (أم أو هاتف متطابق)'
    if sig and sig <= {'DIFF', 'SIB'}:
        return ('SIB', 'إخوة/توائم على الأرجح (نفس الأم، قومي مختلف)') if 'SIB' in sig \
            else ('DIFF', 'أشخاص مختلفون على الأرجح (قومي وأم/ميلاد مختلفة)')
    return 'REVIEW', 'يحتاج مراجعة يدوية'


# ======================================================================
# كتابة ملف Excel
# ======================================================================
HEADERS = ['مجموعة', 'سبب الإدراج', 'التقييم الآلي', 'id', 'الكود', 'رقم القيد', 'الاسم', 'الرقم القومي', 'قومي صالح؟',
           'اسم الأم', 'تاريخ الميلاد', 'الهاتف', 'نوع الدراسة', 'الصف', 'الفصل', 'حالة القيد', 'تاريخ إنشاء السجل',
           'مسكن هذه السنة؟', 'عدد الأقساط', 'المطلوب هذه السنة', 'إجمالي الحساب', 'الخصم', 'عدد الإيصالات (كل السنوات)',
           'المحصل هذه السنة', 'المحصل كل السنوات', 'مديونية قديمة', 'علاقات أخرى']
MONEY_COLS = {'المطلوب هذه السنة', 'إجمالي الحساب', 'الخصم', 'المحصل هذه السنة', 'المحصل كل السنوات', 'مديونية قديمة'}
WIDTHS = {'مجموعة': 8, 'سبب الإدراج': 30, 'التقييم الآلي': 34, 'id': 8, 'الكود': 13, 'رقم القيد': 10, 'الاسم': 32,
          'الرقم القومي': 18, 'قومي صالح؟': 9, 'اسم الأم': 22, 'تاريخ الميلاد': 12, 'الهاتف': 14, 'نوع الدراسة': 10, 'الصف': 16,
          'الفصل': 10, 'حالة القيد': 11, 'تاريخ إنشاء السجل': 13, 'مسكن هذه السنة؟': 9, 'عدد الأقساط': 8,
          'المطلوب هذه السنة': 13, 'إجمالي الحساب': 13, 'الخصم': 10, 'عدد الإيصالات (كل السنوات)': 11,
          'المحصل هذه السنة': 13, 'المحصل كل السنوات': 13,
          'مديونية قديمة': 12, 'علاقات أخرى': 45}


def f(x):
    return float(x) if isinstance(x, Decimal) else x


def build_workbook(ctx):
    """ctx: year, exact, similar, regno, fin(dict by id), rels(dict by id), payments(list), installments(list), threshold"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill('solid', fgColor='1F2937')
    head_font = Font(bold=True, color='FFFFFF')
    band = [PatternFill('solid', fgColor='FFFFFF'), PatternFill('solid', fgColor='EEF2F7')]
    verdict_fill = {'SAME': 'F8B4B4', 'TEMP': 'F8B4B4', 'LIKELY': 'FDE68A', 'REVIEW': 'FDE68A', 'DIFF': 'BBF7D0', 'SIB': 'BBF7D0'}

    def style_sheet(ws, widths):
        ws.sheet_view.rightToLeft = True
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i, h in enumerate([c.value for c in ws[1]], 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
        ws.freeze_panes = 'A2'
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    ws0 = wb.active
    ws0.title = 'ملخص'

    members_all = {}  # id -> (student, [labels])

    def groups_sheet(title, groups, prefix):
        ws = wb.create_sheet(title)
        ws.append(HEADERS)
        for gi, (reason, members) in enumerate(groups, 1):
            label = f'{prefix}{gi}'
            code, verdict = group_verdict(members)
            for s in members:
                members_all.setdefault(s.pk, (s, []))[1].append(label)
                fin = ctx['fin'].get(s.pk, {})
                row = [label, reason, verdict, s.pk, s.student_code or '', s.registration_number or '',
                       full_name(s), clean_nid(s.national_id), 'نعم' if valid_nid(s.national_id) else 'لا (وهمي/ناقص)', s.mother_name or '', s.date_of_birth, s.phone or '',
                       s.get_study_type_display() if hasattr(s, 'get_study_type_display') else '',
                       str(s.grade or ''), str(s.classroom or ''),
                       s.get_enrollment_status_display() if hasattr(s, 'get_enrollment_status_display') else '',
                       s.created_at.date() if getattr(s, 'created_at', None) else '',
                       'نعم' if fin.get('n_inst') else 'لا', fin.get('n_inst', 0), f(fin.get('required', 0)),
                       f(fin.get('acc_fees', 0)), f(fin.get('discount', 0)), fin.get('n_pay', 0),
                       f(fin.get('paid', 0)), f(fin.get('paid_all', 0)), f(s.previous_debt or 0), ctx['rels'].get(s.pk, '')]
                ws.append(row)
                r = ws.max_row
                fill = band[gi % 2]
                for c in ws[r]:
                    c.fill = fill
                    c.alignment = Alignment(vertical='center', wrap_text=True)
                ws.cell(r, 3).fill = PatternFill('solid', fgColor=verdict_fill[code])
                for ci, h in enumerate(HEADERS, 1):
                    if h in MONEY_COLS:
                        ws.cell(r, ci).number_format = '#,##0.00'
                ws.cell(r, 5).number_format = ws.cell(r, 8).number_format = '@'
        style_sheet(ws, WIDTHS)
        return ws

    groups_sheet('متطابقة', ctx['exact'], 'م')
    groups_sheet('متشابهة', ctx['similar'], 'ش')
    groups_sheet('رقم القيد المكرر', ctx['regno'], 'ق')

    # ---- المعاملات المالية (الإيصالات) ----
    wsp = wb.create_sheet('المعاملات المالية')
    ph = ['مجموعة', 'id الطالب', 'الاسم', 'رقم الإيصال', 'تاريخ الدفع', 'السنة', 'فئة الإيراد', 'المبلغ', 'ملغي؟',
          'أُغلقت الخزينة؟', 'تسوية فقط؟', 'المحصّل', 'ملاحظات']
    wsp.append(ph)
    for p in ctx['payments']:
        sid = p['student_id']
        s, labels = members_all.get(sid, (None, []))
        wsp.append([' / '.join(labels), sid, full_name(s) if s else '', p['receipt'] or '', p['date'], p.get('year', ''), p['category'] or '',
                    f(p['amount']), 'نعم' if p['cancelled'] else '', 'نعم' if p['closed'] else '',
                    'نعم' if p['settlement'] else '', p['by'] or '', p['notes'] or ''])
        wsp.cell(wsp.max_row, 8).number_format = '#,##0.00'
    style_sheet(wsp, {'مجموعة': 12, 'الاسم': 30, 'رقم الإيصال': 14, 'فئة الإيراد': 22, 'المبلغ': 12, 'ملاحظات': 40,
                      'المحصّل': 14, 'تاريخ الدفع': 12})

    # ---- الأقساط ----
    wsi = wb.create_sheet('الأقساط')
    ih = ['مجموعة', 'id الطالب', 'الاسم', 'السنة', 'رقم القسط', 'تاريخ الاستحقاق', 'المطلوب', 'المدفوع', 'غرامة تأخير', 'الحالة']
    wsi.append(ih)
    for i in ctx['installments']:
        sid = i['student_id']
        s, labels = members_all.get(sid, (None, []))
        wsi.append([' / '.join(labels), sid, full_name(s) if s else '', i['year'], i['no'], i['due'],
                    f(i['amount']), f(i['paid']), f(i['late']), i['status']])
        for c in (7, 8, 9):
            wsi.cell(wsi.max_row, c).number_format = '#,##0.00'
    style_sheet(wsi, {'مجموعة': 12, 'الاسم': 30, 'المطلوب': 12, 'المدفوع': 12, 'تاريخ الاستحقاق': 14})

    # ---- الملخص ----
    lines = [
        ('السنة الدراسية', str(ctx['year'])),
        ('عدد الطلاب المفحوصين', ctx['scanned']),
        ('مجموعات الأسماء المتطابقة (ورقة "متطابقة")', len(ctx['exact'])),
        ('أزواج الأسماء المتشابهة (ورقة "متشابهة") – حد التشابه', f"{len(ctx['similar'])}  (≥ {ctx['threshold']:.0%})"),
        ('مجموعات رقم القيد المكرر', len(ctx['regno'])),
        ('عدد الإيصالات في ورقة المعاملات', len(ctx['payments'])),
        ('', ''),
        ('طريقة القراءة', 'كل "مجموعة" هي طلاب يُشتبه في تكرارهم. الرمز: م = متطابق، ش = متشابه، ق = رقم قيد مكرر.'),
        ('لون التقييم الآلي', 'أحمر = نفس الشخص أو سجل مؤقت (قومي غير صالح) | أصفر = يحتاج مراجعة | أخضر = أشخاص مختلفون أو إخوة/توائم'),
        ('تنبيه', 'التقييم الآلي استرشادي فقط. قارن الرقم القومي واسم الأم وتاريخ الميلاد وحساباتهم قبل أي إجراء.'),
        ('تنبيه', 'اسمان بنفس الاسم لكن برقم قومي مختلف غالباً شخصان مختلفان (أسماء شائعة). لا يُحذف أي سجل له معاملات.'),
    ]
    ws0.append(['البند', 'القيمة'])
    for a, b in lines:
        ws0.append([a, b])
    style_sheet(ws0, {'البند': 52, 'القيمة': 95})
    for row in ws0.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='center')
    return wb


# ======================================================================
# الأمر
# ======================================================================
class Command(BaseCommand):
    help = 'تصدير Excel بالأسماء المتطابقة والمتشابهة وحساباتها المالية للمراجعة (لا يعدّل بيانات)'

    def add_arguments(self, parser):
        parser.add_argument('--year-id', type=int, help='السنة الدراسية (الافتراضي: النشطة)')
        parser.add_argument('--workers-only', action='store_true', help='الطلاب العمال فقط')
        parser.add_argument('--threshold', type=float, default=0.88, help='حد التشابه 0-1 (الافتراضي 0.88)')
        parser.add_argument('--out', type=str, help='مسار ملف الإخراج .xlsx')

    def handle(self, *args, **o):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            raise CommandError('مكتبة openpyxl غير مثبتة. شغّل:  pip install openpyxl')

        from django.core.exceptions import ObjectDoesNotExist
        from django.db.models import Count, Sum
        from finance.models import AcademicYear, Payment, StudentAccount, StudentInstallment
        from students.models import Student

        year = (AcademicYear.objects.filter(id=o['year_id']).first() if o['year_id']
                else AcademicYear.objects.filter(is_active=True).first())
        if not year:
            raise CommandError('لم يتم العثور على السنة الدراسية.')

        qs = Student.objects.filter(academic_year=year).select_related('grade', 'classroom')
        if o['workers_only']:
            qs = qs.filter(study_type='Workers')
        students = list(qs)
        self.stdout.write(f'تم تحميل {len(students)} طالب. جاري البحث عن التشابه...')

        exact = exact_groups(students)
        similar = similar_pairs(students, o['threshold'])
        regno = registration_groups(students)
        self.stdout.write(f'متطابقة: {len(exact)} | متشابهة: {len(similar)} | رقم قيد مكرر: {len(regno)}')

        by_id = {s.pk: s for _, ms in (exact + similar + regno) for s in ms}
        ids = list(by_id)
        if not ids:
            self.stdout.write(self.style.SUCCESS('لا توجد أسماء متشابهة أو مكررة.'))
            return

        # ---------- الحسابات ----------
        fin = defaultdict(dict)
        payments, installments = [], []
        for i in range(0, len(ids), 500):
            part = ids[i:i + 500]
            for sid, n, t in (StudentInstallment.objects.filter(student_id__in=part, academic_year=year)
                              .values('student').annotate(n=Count('id'), t=Sum('amount_due'))
                              .values_list('student', 'n', 't')):
                fin[sid]['n_inst'], fin[sid]['required'] = n, t or Decimal('0')
            for sid, t, d in (StudentAccount.objects.filter(student_id__in=part, academic_year=year)
                              .values('student').annotate(t=Sum('total_fees'), d=Sum('discount'))
                              .values_list('student', 't', 'd')):
                fin[sid]['acc_fees'], fin[sid]['discount'] = t or Decimal('0'), d or Decimal('0')
            for sid, n in (Payment.objects.filter(student_id__in=part).values('student')
                           .annotate(n=Count('id')).values_list('student', 'n')):
                fin[sid]['n_pay'] = n
            for sid, t in (Payment.objects.filter(student_id__in=part, is_cancelled=False, academic_year=year)
                           .values('student').annotate(t=Sum('amount_paid')).values_list('student', 't')):
                fin[sid]['paid'] = t or Decimal('0')
            for sid, t in (Payment.objects.filter(student_id__in=part, is_cancelled=False).values('student')
                           .annotate(t=Sum('amount_paid')).values_list('student', 't')):
                fin[sid]['paid_all'] = t or Decimal('0')
            for p in (Payment.objects.filter(student_id__in=part)
                      .select_related('revenue_category', 'collected_by', 'academic_year').order_by('student_id', 'payment_date', 'id')):
                payments.append({
                    'student_id': p.student_id, 'receipt': p.receipt_number, 'date': p.payment_date,
                    'year': str(p.academic_year) if p.academic_year_id else '',
                    'category': p.revenue_category.name if p.revenue_category_id else '',
                    'amount': p.amount_paid, 'cancelled': p.is_cancelled, 'closed': p.is_closed,
                    'settlement': p.is_settlement_only,
                    'by': p.collected_by.username if p.collected_by_id else '', 'notes': p.notes})
            for x in (StudentInstallment.objects.filter(student_id__in=part)
                      .select_related('academic_year').order_by('student_id', 'academic_year_id', 'due_date')):
                installments.append({
                    'student_id': x.student_id, 'year': str(x.academic_year) if x.academic_year_id else '',
                    'no': getattr(x, 'installment_number', ''), 'due': x.due_date, 'amount': x.amount_due,
                    'paid': x.paid_amount, 'late': getattr(x, 'late_fee', 0) or 0, 'status': x.status})

        # ---------- العلاقات الأخرى (دفتر اليومية، الحضور، ...) ----------
        rels = {}
        for s in by_id.values():
            parts = []
            for rel in Student._meta.related_objects:
                name = rel.get_accessor_name()
                try:
                    obj = getattr(s, name)
                except ObjectDoesNotExist:
                    continue
                n = 1 if rel.one_to_one else obj.count()
                if n:
                    parts.append(f'{name}:{n}')
            rels[s.pk] = '، '.join(parts)

        wb = build_workbook({
            'year': year, 'scanned': len(students), 'threshold': o['threshold'], 'exact': exact, 'similar': similar,
            'regno': regno, 'fin': fin, 'rels': rels, 'payments': payments, 'installments': installments})

        out = o['out'] or os.path.join(os.path.expanduser('~'), f"similar_students_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(out)
        self.stdout.write(self.style.SUCCESS(f'\nتم إنشاء الملف: {out}'))
        self.stdout.write('نزّله من صفحة Files في PythonAnywhere وافتحه في Excel.')
